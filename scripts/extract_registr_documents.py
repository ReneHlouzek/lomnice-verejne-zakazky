#!/usr/bin/env python3
"""Download and text-extract attachments exposed by the official Contract Register."""
from __future__ import annotations
import hashlib, io, json, time, re, unicodedata, subprocess
from urllib.parse import unquote
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
import requests
from pypdf import PdfReader
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/"data"/"sources"/"registr-smluv"
OUT=ROOT/"data"/"documents"/"registr-smluv"
MANIFEST=OUT/"manifest.json"
TIMEOUT=20
MAX_DOC_BYTES=25*1024*1024
WORKERS=6
UA="Lomnice-Verejne-Zakazky/1.0 (public-data-archive)"

def extract_legacy_doc(target: Path, temp_dir: str) -> str:
    """Extract a legacy .doc file, falling back to LibreOffice when antiword is absent."""
    import subprocess
    try:
        cp = subprocess.run(["antiword", str(target)], check=True, capture_output=True, text=True)
        return cp.stdout.strip()
    except FileNotFoundError:
        subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "txt:Text", "--outdir", temp_dir, str(target)],
            check=True, capture_output=True, text=True,
        )
        out = target.with_suffix(".txt")
        return out.read_text(encoding="utf-8", errors="replace").strip() if out.exists() else ""


def extract(url:str):
    suffix=url.lower().split("?",1)[0].rsplit(".",1)[-1]
    if suffix in ("docx", "doc", "odt"):
        import subprocess, tempfile
        with tempfile.TemporaryDirectory() as td:
            target=Path(td)/("input."+suffix)
            try:
                rr=requests.get(url,timeout=TIMEOUT,headers={"User-Agent":UA,"Accept":"*/*"})
                rr.raise_for_status()
                data=rr.content
            except requests.RequestException:
                cp=subprocess.run(["curl","--fail","--location","--http1.1","--retry","2","--retry-delay","1","--connect-timeout","8","--max-time","25","-A",UA,url],check=True,capture_output=True)
                data=cp.stdout
            if len(data)>MAX_DOC_BYTES: raise ValueError(f"document too large: {len(data)} bytes")
            meta={"url":(rr.url if "rr" in locals() else url),"sha256":hashlib.sha256(data).hexdigest(),"bytes":len(data),"content_type":(rr.headers.get("content-type","") if "rr" in locals() else "")}
            target.write_bytes(data)
            if suffix=="docx":
                doc=Document(str(target))
                parts=[p.text for p in doc.paragraphs]
                for table in doc.tables:
                    parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
                return meta,"\n".join(parts).strip()
            if suffix=="doc":
                return meta,extract_legacy_doc(target,td)
            subprocess.run(["libreoffice","--headless","--convert-to","txt:Text","--outdir",td,str(target)],check=True,capture_output=True,text=True)
            out=target.with_suffix(".txt")
            return meta,(out.read_text(encoding="utf-8",errors="replace") if out.exists() else "").strip()
    try:
        r=requests.get(url,timeout=TIMEOUT,headers={"User-Agent":UA,"Accept":"application/pdf,*/*"})
        r.raise_for_status()
        data=r.content
    except requests.RequestException:
        import subprocess
        cp=subprocess.run(
            ["curl","--fail","--location","--http1.1","--retry","2","--retry-delay","1",
             "--connect-timeout","8","--max-time","25","-A",UA,
             "-H","Accept: application/pdf,*/*",url],
            check=True,capture_output=True
        )
        data=cp.stdout
    if len(data)>MAX_DOC_BYTES: raise ValueError(f"document too large: {len(data)} bytes")
    meta={"url":(r.url if "r" in locals() else url),"sha256":hashlib.sha256(data).hexdigest(),"bytes":len(data),
          "content_type":(r.headers.get("content-type","") if "r" in locals() else "")}
    if not data.startswith(b"%PDF"): return meta,None
    reader=PdfReader(io.BytesIO(data))
    parts=[]
    for page in reader.pages:
        try: parts.append(page.extract_text() or "")
        except Exception: parts.append("")
    return meta,"\n\n".join(parts).strip()

def normalized_words(value:str)->set[str]:
    value=unicodedata.normalize("NFKD", unquote(value or ""))
    value="".join(ch for ch in value if not unicodedata.combining(ch)).lower()
    return set(re.findall(r"[a-z]{4,}", value))

def filename_content_mismatch(name:str, text:str|None)->bool:
    """Flag likely wrong attachment payloads when a distinctive filename name is absent from extracted text."""
    decoded=unquote(name or "")
    stem=decoded.rsplit("/",1)[-1].rsplit(".",1)[0]
    words=normalized_words(stem)
    generic={
        "dodatek","smlouva","smlouvy","priloha","prilohy","dokument","scan",
        "final","verze","navrh","podpis","podepsano","signed","document",
        "contract","agreement","attachment","copy","kopie","strana","strany",
        "cislo","castka","rozpocet","projekt","mesto","obec","kraj"
    }
    distinctive=words-generic
    if not distinctive or not text:
        return False
    body_words=normalized_words(text)
    # A single meaningful filename token is enough to establish overlap. If
    # none appears in the extracted body, retain the text but require review.
    return not bool(distinctive & body_words)

def process(path:Path):
    obj=json.loads(path.read_text(encoding="utf-8"))
    rec=obj.get("record",obj)
    sid=str(rec.get("source_id") or path.stem)
    raw=rec.get("raw") or {}
    attachments=rec.get("attachments") or raw.get("attachments") or []
    seen=set()
    result={"source":"registr-smluv","source_id":sid,"title":rec.get("title"),
            "page_url":rec.get("source_url"),"documents":[],"status":"ok"}
    for i,a in enumerate(attachments,1):
        url=a.get("url") if isinstance(a,dict) else str(a)
        name=a.get("name","") if isinstance(a,dict) else ""
        if not url or url in seen or url.lower().split("?",1)[0].rsplit(".",1)[-1] not in ("pdf","doc","docx","odt"):
            continue
        seen.add(url)
        clean_name=name
        if "https://smlouvy.gov.cz/smlouva/soubor/" in clean_name:
            clean_name=clean_name.split("https://smlouvy.gov.cz/smlouva/soubor/",1)[0]
        clean_name=clean_name.strip()
        item={"index":i,"name":clean_name,"url":url}
        existing=OUT/sid/f"{i:03d}.txt"
        if existing.exists():
            text=existing.read_text(encoding="utf-8",errors="replace")
            item["text_available"]=bool(text)
            item["text_chars"]=len(text)
            item["text_file"]=f"data/documents/registr-smluv/{sid}/{i:03d}.txt"
            if filename_content_mismatch(clean_name, text):
                item["status"]="content_mismatch_review"
                item["content_check"]="distinctive_filename_terms_absent_from_text"
            else:
                item["status"]="retained_existing_text"
                item["content_check"]="no_filename_mismatch_detected"
            result["documents"].append(item)
            continue
        try:
            meta,text=extract(url)
            item.update(meta)
            item["text_available"]=bool(text)
            item["text_chars"]=len(text or "")
            if filename_content_mismatch(clean_name, text):
                item["status"]="content_mismatch_review"
                item["content_check"]="distinctive_filename_terms_absent_from_text"
            elif text:
                item["content_check"]="no_filename_mismatch_detected"
            if text:
                outdir=OUT/sid
                outdir.mkdir(parents=True,exist_ok=True)
                item["text_file"]=f"data/documents/registr-smluv/{sid}/{i:03d}.txt"
                (outdir/f"{i:03d}.txt").write_text(text,encoding="utf-8")
        except Exception as exc:
            item["status"]="download_error"
            item["error"]=str(exc)
        result["documents"].append(item)
    return result

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    previous={}
    if MANIFEST.exists():
        try:
            old=json.loads(MANIFEST.read_text(encoding="utf-8"))
            for rec in old.get("records",[]):
                previous[str(rec.get("source_id"))]=rec
        except Exception:
            previous={}
    # The importer can retain both normalized and raw snapshots. Process each
    # official contract only once by source_id.
    all_files=sorted(SRC.glob("*.json"))
    by_id={}
    for p in all_files:
        try:
            o=json.loads(p.read_text(encoding="utf-8"))
            rec=o.get("record",o)
            sid=str(rec.get("source_id") or p.stem)
            by_id.setdefault(sid,p)
        except Exception:
            continue
    files=list(by_id.values())
    results=[]
    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        futures={pool.submit(process,p):p for p in files}
        for fut in as_completed(futures):
            p=futures[fut]
            try: results.append(fut.result())
            except Exception as exc: results.append({"source":"registr-smluv","source_id":p.stem,"documents":[],"status":"worker_error","error":str(exc)})
    results.sort(key=lambda x:x.get("source_id",""))
    for result in results:
        old=previous.get(str(result.get("source_id")))
        if not old:
            continue
        old_docs={d.get("url"):d for d in old.get("documents",[]) if d.get("url")}
        for doc in result.get("documents",[]):
            old_doc=old_docs.get(doc.get("url"))
            if old_doc and not doc.get("sha256") and not doc.get("content_check"):
                retained={k:v for k,v in old_doc.items() if k not in ("status","error")}
                retained["status"]="retained_from_previous_run"
                doc.clear()
                doc.update(retained)
    def get_attachments(o):
        rec=o.get("record",o)
        return rec.get("attachments") or (rec.get("raw") or {}).get("attachments") or []
    declared=sum(len(get_attachments(json.loads(p.read_text(encoding="utf-8")))) for p in files)
    discovered=sum(len(x.get("documents",[])) for x in results)
    downloaded=sum(1 for x in results for d in x.get("documents",[]) if d.get("sha256"))
    extracted=sum(1 for x in results for d in x.get("documents",[]) if d.get("text_available"))
    MANIFEST.write_text(json.dumps({
        "updated_at":time.strftime("%Y-%m-%dT%H:%M:%SZ",time.gmtime()),
        "records_checked":len(results),"attachments_declared":declared,
        "documents_discovered":discovered,"documents_downloaded":downloaded,
        "documents_text_extracted":extracted,"records":results,
        "note":"Original binaries are not committed; extracted text and document metadata are retained."
    },ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(f"RS documents: records={len(results)} declared={declared} discovered={discovered} downloaded={downloaded} text={extracted}")

if __name__=="__main__":
    main()
